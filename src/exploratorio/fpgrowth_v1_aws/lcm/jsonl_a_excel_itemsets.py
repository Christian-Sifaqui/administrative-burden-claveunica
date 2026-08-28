from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


INPUT_JSONL = "final_output_lcm_closed_traducido_2024_2025.jsonl"
OUTPUT_XLSX = "itemsets_lcm_2024_2025.xlsx"


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def normalize_record(obj: Any) -> tuple[list[str], int]:
    """
    Espera líneas JSON con forma:
      [[item1, item2, ...], support_abs]
    """
    if not isinstance(obj, list) or len(obj) != 2:
        raise ValueError(f"Formato inesperado de línea JSON: {obj!r}")

    items, support_abs = obj

    if not isinstance(items, list):
        raise ValueError(f"La primera posición debe ser lista de items: {obj!r}")

    items = [str(x) for x in items]

    try:
        support_abs = int(support_abs)
    except Exception as exc:
        raise ValueError(f"El soporte absoluto no pudo convertirse a int: {obj!r}") from exc

    return items, support_abs


def autosize_worksheet(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_idx = col_cells[0].column
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        adjusted = min(max_len + 2, 80)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER


def create_sheet_with_rows(wb: Workbook, title: str, rows: list[tuple[list[str], int]], size: int) -> None:
    ws = wb.create_sheet(title=title)

    headers = ["tamano_itemset"] + [f"item_{i}" for i in range(1, size + 1)] + ["support_abs"]
    ws.append(headers)

    for items, support_abs in rows:
        row = [len(items)] + items + [support_abs]
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    style_header(ws)
    autosize_worksheet(ws)


def main() -> None:
    input_path = Path(INPUT_JSONL)
    output_path = Path(OUTPUT_XLSX)

    if not input_path.exists():
        raise FileNotFoundError(f"No existe el archivo de entrada: {input_path}")

    size_1: list[tuple[list[str], int]] = []
    size_2: list[tuple[list[str], int]] = []
    size_3: list[tuple[list[str], int]] = []
    skipped = 0
    total = 0

    with input_path.open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue

            total += 1
            try:
                obj = json.loads(line)
                items, support_abs = normalize_record(obj)
            except Exception as exc:
                raise ValueError(f"Error en línea {line_no}: {exc}") from exc

            n = len(items)
            if n == 1:
                size_1.append((items, support_abs))
            elif n == 2:
                size_2.append((items, support_abs))
            elif n == 3:
                size_3.append((items, support_abs))
            else:
                skipped += 1

    # Orden descendente por soporte, y luego alfabético por items
    size_1.sort(key=lambda x: (-x[1], x[0]))
    size_2.sort(key=lambda x: (-x[1], x[0]))
    size_3.sort(key=lambda x: (-x[1], x[0]))

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "resumen"
    ws0.append(["hoja", "n_filas"])
    ws0.append(["conjuntos_1", len(size_1)])
    ws0.append(["conjuntos_2", len(size_2)])
    ws0.append(["conjuntos_3", len(size_3)])
    ws0.append(["omitidos_otro_tamano", skipped])
    ws0.append(["total_lineas_leidas", total])
    ws0.freeze_panes = "A2"
    ws0.auto_filter.ref = ws0.dimensions
    style_header(ws0)
    autosize_worksheet(ws0)

    create_sheet_with_rows(wb, "conjuntos_1", size_1, size=1)
    create_sheet_with_rows(wb, "conjuntos_2", size_2, size=2)
    create_sheet_with_rows(wb, "conjuntos_3", size_3, size=3)

    wb.save(output_path)
    print(f"Excel generado: {output_path.resolve()}")
    print(f"Filas tamaño 1: {len(size_1):,}")
    print(f"Filas tamaño 2: {len(size_2):,}")
    print(f"Filas tamaño 3: {len(size_3):,}")
    print(f"Omitidos por tamaño distinto de 1, 2 o 3: {skipped:,}")


if __name__ == "__main__":
    main()
